"""Frozen-source assertions and SYNTHETIC in-memory mutations; never live data."""

import copy
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

import openpyxl

from importer.__main__ import MAX_BYTES, atomic_json, check_archive, download, main
from importer.workbook import DIVISIONS, ImportFailure, parse_workbook, validate_data

FIXTURE = Path(__file__).parent / "fixtures/source-2026-10-13.xlsx"
# Deterministic SYNTHETIC timestamp, used only by tests, not published as a download.
TEST_TIME = "2026-10-13T00:00:00Z"


def synthetic_workbook(changes: dict) -> bytes:
    """Apply isolated cells in memory; the frozen file is never saved/modified."""
    workbook = openpyxl.load_workbook(FIXTURE)
    for sheet, cells in changes.items():
        for coordinate, value in cells.items():
            workbook[sheet][coordinate] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class ImporterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.payload = FIXTURE.read_bytes()
        cls.data = parse_workbook(cls.payload, TEST_TIME, "frozen")

    def test_frozen_counts_and_empty_features(self) -> None:
        data = self.data
        self.assertEqual(len(data["teams"]), 12)
        self.assertEqual(len(data["matches"]), 56)
        self.assertEqual(len({match["id"] for match in data["matches"]}), 56)
        self.assertEqual([sum(match["phase"] == phase for match in data["matches"]) for phase in ("first", "second", "finals")], [30, 12, 14])
        self.assertEqual(len(data["carryovers"]), 6)
        self.assertTrue(all(match["resultState"] == "none" for match in data["matches"]))
        self.assertEqual(len(data["sotg"]), 112)
        self.assertTrue(all(value is None for evaluation in data["sotg"] for value in evaluation["categories"]))
        self.assertEqual(data["players"], [])
        self.assertFalse(any(data["availability"].values()))
        self.assertEqual(data["finalClassification"], [])
        self.assertEqual(data["sourceReportedAt"], "2026-10-11T16:15:23+02:00")

    def test_five_first_phase_games_per_team(self) -> None:
        for team in self.data["teams"]:
            self.assertEqual(sum(match["phase"] == "first" and team["id"] in (match["a"]["teamId"], match["b"]["teamId"]) for match in self.data["matches"]), 5)

    def test_source_form_links_preserve_direction_and_flag_wrong_prefills(self) -> None:
        evaluations = self.data["sotg"]
        self.assertEqual(sum(evaluation["formUrl"] is not None for evaluation in evaluations), 112)
        for evaluation in evaluations:
            query = parse_qs(urlsplit(evaluation["formUrl"]).query, keep_blank_values=True)
            self.assertEqual(query["entry.1623410474"], [evaluation["matchId"]])
        first = [evaluation for evaluation in evaluations if evaluation["matchId"] == "X_1"]
        self.assertNotEqual(first[0]["formUrl"], first[1]["formUrl"])
        self.assertTrue(all(evaluation["formTeamsMatch"] for evaluation in first))
        second = [evaluation for evaluation in evaluations if evaluation["matchId"] == "X_2"]
        self.assertTrue(all(evaluation["formTeamsMatch"] is False for evaluation in second))
        final = [evaluation for evaluation in evaluations if evaluation["matchId"] == "F1"]
        self.assertEqual(final[0]["formUrl"], final[1]["formUrl"])
        self.assertTrue(all(evaluation["formTeamsMatch"] is None for evaluation in final))

    def test_synthetic_form_link_missing_unsafe_or_wrong_match(self) -> None:
        for url in ["javascript:alert(1)", "https://example.com/viewform", self.data["sotg"][0]["formUrl"].replace("entry.1623410474=X_1", "entry.1623410474=X_2")]:
            data = copy.deepcopy(self.data)
            data["sotg"][0]["formUrl"] = url
            with self.assertRaises(ImportFailure):
                validate_data(data)
        book = openpyxl.load_workbook(io.BytesIO(self.payload))
        sheet = book["Valutazioni SOTG"]
        sheet["E2"].hyperlink = None
        buffer = io.BytesIO()
        book.save(buffer)
        with self.assertRaisesRegex(ImportFailure, "link label without hyperlink"):
            parse_workbook(buffer.getvalue(), TEST_TIME)
        sheet["E2"] = None
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        data = parse_workbook(buffer.getvalue(), TEST_TIME)
        self.assertIsNone(data["sotg"][0]["formUrl"])

    def test_reference_fixtures_and_merges(self) -> None:
        games = {match["id"]: match for match in self.data["matches"]}
        names = {team["id"]: team["name"] for team in self.data["teams"]}
        first = games["X_1"]
        self.assertEqual((names[first["a"]["teamId"]], names[first["b"]["teamId"]]), ("BFD Redshot", "VOLADORA"))
        for identifier, expected in {"X_1": ("PARMA", "2026-10-25", "09:00", None),
                                     "X_7": ("PADOVA", "2026-11-07", "10:30", "campo 2"),
                                     "Y_13": ("PADOVA", "2026-11-08", "10:30", "campo 1"),
                                     "F1": ("PADOVA", "2026-11-22", "15:00", "campo 1")}.items():
            match = games[identifier]
            self.assertEqual(tuple(match[key] for key in ("city", "date", "start", "field")), expected)
        self.assertEqual((games["F1"]["a"]["slot"], games["F1"]["b"]["slot"]), ("WS1", "WS2"))
        self.assertIsNone(games["X_6"]["field"])
        self.assertEqual(games["X_2"]["city"], "PARMA")
        self.assertTrue(all(entry["teamId"] is None for pool in self.data["pools"] if pool["phase"] == "second" for entry in pool["entries"]))

    def test_synthetic_zero_and_partial_result(self) -> None:
        data = parse_workbook(synthetic_workbook({"Partite": {"L5": 0, "N5": 15, "L6": 0}}), TEST_TIME)
        games = {match["id"]: match for match in data["matches"]}
        self.assertEqual(games["X_1"]["score"], {"a": 0, "b": 15})
        self.assertEqual(games["X_1"]["resultState"], "complete")
        self.assertEqual(games["X_2"]["resultState"], "partial")
        self.assertIsNone(games["X_2"]["score"]["b"])
        self.assertEqual(games["X_3"]["resultState"], "none")

    def test_synthetic_report_marker_preserves_link_and_finished_score(self) -> None:
        book = openpyxl.load_workbook(io.BytesIO(self.payload))
        sheet = book["Partite"]
        sheet["Q5"] = "report"
        url = "https://docs.google.com/spreadsheets/d/SYNTHETIC_REPORT/edit?gid=123#gid=123"
        sheet["Q5"].hyperlink = url
        sheet["L5"], sheet["N5"] = 0, 15
        buffer = io.BytesIO()
        book.save(buffer)
        data = parse_workbook(buffer.getvalue(), TEST_TIME)
        self.assertEqual(len(data["carryovers"]), 6)
        self.assertEqual(data["matches"][0]["reportUrl"], url)
        self.assertEqual(data["matches"][0]["score"], {"a": 0, "b": 15})
        self.assertEqual(data["matches"][0]["resultState"], "complete")
        for unsafe in (None, "javascript:alert(1)", "https://example.com/report"):
            sheet["Q5"].hyperlink = unsafe
            buffer = io.BytesIO()
            book.save(buffer)
            with self.subTest(unsafe=unsafe), self.assertRaises(ImportFailure):
                parse_workbook(buffer.getvalue(), TEST_TIME)
        book.close()

    def test_synthetic_unnamed_player_totals_are_preserved(self) -> None:
        data = parse_workbook(synthetic_workbook({"Statistiche": {
            "A3": "BFD Redshot", "N3": 3, "O3": 0, "P3": 3}}), TEST_TIME)
        self.assertIsNone(data["players"][0]["surname"])
        self.assertIsNone(data["players"][0]["givenName"])
        self.assertEqual(data["players"][0]["total"], 3)
        with self.assertRaises(ImportFailure):
            parse_workbook(synthetic_workbook({"Statistiche": {"A3": "BFD Redshot", "B3": 99}}), TEST_TIME)

    def test_synthetic_opposite_sotg_directions(self) -> None:
        cells = {f"{column}{row}": 0 for row in (2, 3) for column in "GHIJKL"}
        cells["M2"] = "SYNTHETIC <script>test</script>"
        data = parse_workbook(synthetic_workbook({"Valutazioni SOTG": cells}), TEST_TIME)
        evaluations = [evaluation for evaluation in data["sotg"] if evaluation["matchId"] == "X_1"]
        self.assertEqual(len(evaluations), 2)
        self.assertEqual(evaluations[0]["evaluatingTeamId"], evaluations[1]["evaluatedTeamId"])
        self.assertEqual(evaluations[0]["categories"], [0] * 5)
        self.assertTrue(data["availability"]["sotg"])

    def test_synthetic_carryover_not_new_match(self) -> None:
        data = parse_workbook(synthetic_workbook({"Partite": {"L41": 0, "N41": 15}}), TEST_TIME)
        self.assertEqual(len(data["matches"]), 56)
        self.assertEqual(data["carryovers"][0]["score"], {"a": 0, "b": 15})

    def test_synthetic_published_players_are_team_list(self) -> None:
        data = parse_workbook(synthetic_workbook({"Statistiche": {"A3": "BFD Redshot", "B3": "Test", "C3": "Synthetic", "N3": 0, "O3": 2, "P3": 2, "D3": "0 | 2"}}), TEST_TIME)
        self.assertEqual(len(data["players"]), 1)
        self.assertEqual(data["players"][0]["assists"], 0)
        self.assertTrue(data["availability"]["teamPlayers"])
        self.assertFalse(data["availability"]["matchPlayers"])
        self.assertFalse(data["availability"]["roster"])

    def test_synthetic_bad_headers_duplicate_ids_and_unknown_names(self) -> None:
        for changes in ({"Partite": {"D4": "changed"}}, {"Partite": {"I6": "X_1"}},
                        {"Partite": {"K5": "Unknown"}}, {"Partite": {"I5": None}},
                        {"Partite": {"Q5": "annullata"}}):
            with self.subTest(changes=changes), self.assertRaises(ImportFailure):
                parse_workbook(synthetic_workbook(changes), TEST_TIME)

    def test_synthetic_bad_scores(self) -> None:
        for value in (-1, "15", float('inf'), 1.2):
            data = copy.deepcopy(self.data)
            data["matches"][0]["score"]["a"] = value
            with self.subTest(value=value), self.assertRaises(ImportFailure):
                validate_data(data)

    def test_synthetic_missing_cached_formula(self) -> None:
        with self.assertRaisesRegex(ImportFailure, "formula has no cached result"):
            parse_workbook(synthetic_workbook({"Partite": {"L5": "=1+1"}}), TEST_TIME)

    def test_synthetic_qualifier_not_guessed(self) -> None:
        match = next(match for match in self.data["matches"] if match["id"] == "P_1")
        self.assertEqual(match["a"], {"slot": "1X", "teamId": None})
        self.assertEqual(self.data["teams"][0]["slot"], "X1")

    def test_synthetic_final_placement_needs_published_team(self) -> None:
        data = parse_workbook(synthetic_workbook({"Classifiche e format": {"O20": "BFD Redshot"}}), TEST_TIME)
        self.assertEqual(data["finalClassification"], [{"position": 1, "teamId": "x1", "provenance": {"sheet": "Classifiche e format", "range": "N20:O20"}}])

    def test_synthetic_invalid_evaluation_direction_and_sum(self) -> None:
        for changes in ({"Valutazioni SOTG": {"C2": "VOLADORA", "D2": "VOLADORA"}},
                        {"Valutazioni SOTG": {"G2": 1, "H2": 1, "I2": 1, "J2": 1, "K2": 1, "L2": 0}},
                        {"Valutazioni SOTG": {"G2": 5}}):
            with self.subTest(changes=changes), self.assertRaises(ImportFailure):
                parse_workbook(synthetic_workbook(changes), TEST_TIME)

    def test_invalid_html_or_oversized_archive(self) -> None:
        for payload in (b"<html>login</html>", b"0" * (MAX_BYTES + 1)):
            with self.assertRaises(ImportFailure):
                check_archive(payload)

    def test_download_retries_are_bounded(self) -> None:
        with patch("urllib.request.urlopen", side_effect=OSError("synthetic network failure")) as request, patch("time.sleep"):
            with self.assertRaisesRegex(ImportFailure, "3 attempts"):
                download()
            self.assertEqual(request.call_count, 3)

    def test_failed_import_preserves_previous_file(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            source, output = directory / "bad.xlsx", directory / "data.json"
            source.write_bytes(b"<html>synthetic error</html>")
            output.write_text('previous valid snapshot')
            process = subprocess.run([sys.executable, "-m", "importer", "--input", str(source), "--fetched-at", TEST_TIME, "--output", str(output)], capture_output=True)
            self.assertEqual(process.returncode, 1)
            self.assertEqual(output.read_text(), 'previous valid snapshot')
            self.assertIn(b"previous output preserved", process.stderr)

    def test_atomic_output_has_no_nan(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "data.json"
            atomic_json(output, self.data)
            self.assertEqual(json.loads(output.read_text()), self.data)


class MixedSerieBTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.payload = (FIXTURE.parent / "source-mixed-b-2026-10-15.xlsx").read_bytes()
        cls.data = parse_workbook(cls.payload, TEST_TIME, "frozen", division_id="mixed-b")

    def test_frozen_division_structure_and_references(self) -> None:
        data = self.data
        self.assertEqual(data["division"], "mixed-b")
        self.assertEqual(data["sourceUrl"], DIVISIONS["mixed-b"]["sourceUrl"])
        self.assertEqual((len(data["teams"]), len(data["matches"]), len(data["pools"])), (18, 63, 9))
        self.assertEqual([sum(match["phase"] == phase for match in data["matches"]) for phase in ("first", "second", "finals")], [45, 18, 0])
        self.assertEqual(data["carryovers"], [])
        self.assertEqual(data["finalClassification"], [])
        self.assertEqual(data["players"], [])
        self.assertFalse(any(data["availability"].values()))
        self.assertEqual(data["sourceReportedAt"], "2026-10-11T16:15:21+02:00")
        self.assertEqual([(pool["id"], len(pool["entries"])) for pool in data["pools"][:3]], [("X", 6), ("Y", 6), ("Z", 6)])
        self.assertTrue(all(len(pool["entries"]) == 3 for pool in data["pools"][3:]))
        for team in data["teams"]:
            self.assertEqual(sum(match["phase"] == "first" and team["id"] in (match["a"]["teamId"], match["b"]["teamId"]) for match in data["matches"]), 5)
        matches = {match["id"]: match for match in data["matches"]}
        self.assertEqual((matches["X_1"]["a"]["teamId"], matches["X_1"]["b"]["teamId"]), ("x3", "x4"))
        self.assertEqual((matches["Z_1"]["a"]["teamId"], matches["Z_1"]["b"]["teamId"]), ("z1", "z5"))
        self.assertEqual(matches["P_18"]["pool"], "power-6")
        self.assertEqual(matches["P_18"]["b"], {"slot": "6Z", "teamId": None})
        self.assertTrue(any("Non sono previste fasi a eliminazione" in rule["text"] for rule in data["rules"]))

    def test_form_links_keep_the_correct_division_and_direction(self) -> None:
        self.assertEqual(len(self.data["sotg"]), 126)
        for evaluation in self.data["sotg"]:
            query = parse_qs(urlsplit(evaluation["formUrl"]).query, keep_blank_values=True)
            self.assertEqual(query["entry.1017677778"], ["Mixed Serie B"])
            self.assertEqual(query["entry.1623410474"], [evaluation["matchId"]])
            self.assertIsNot(evaluation["formTeamsMatch"], False)

    def test_synthetic_wrong_division_or_source_fails(self) -> None:
        with self.assertRaisesRegex(ImportFailure, "expected header"):
            parse_workbook(self.payload, TEST_TIME)
        changed = copy.deepcopy(self.data)
        changed["sourceUrl"] = DIVISIONS["mixed-a"]["sourceUrl"]
        with self.assertRaisesRegex(ImportFailure, "Division/source mismatch"):
            validate_data(changed)

    def test_synthetic_b_header_or_removed_match_fails(self) -> None:
        for sheet, cell, value in [("Classifiche e format", "H51", "UNKNOWN"), ("Partite", "I39", None)]:
            book = openpyxl.load_workbook(io.BytesIO(self.payload))
            book[sheet][cell] = value
            buffer = io.BytesIO()
            book.save(buffer)
            book.close()
            with self.subTest(cell=cell), self.assertRaises(ImportFailure):
                parse_workbook(buffer.getvalue(), TEST_TIME, division_id="mixed-b")

    def test_synthetic_second_division_failure_replaces_neither_output(self) -> None:
        with patch("sys.argv", ["importer", "--division", "all"]), \
                patch("importer.__main__.download", side_effect=[FIXTURE.read_bytes(), b"bad synthetic workbook"]), \
                patch("importer.__main__.atomic_json") as write, patch("sys.stderr", io.StringIO()):
            with self.assertRaises(SystemExit) as raised:
                main()
            self.assertEqual(raised.exception.code, 1)
            write.assert_not_called()


if __name__ == "__main__":
    unittest.main()
