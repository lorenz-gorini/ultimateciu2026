"""Validated XLSX-to-JSON boundary. No formula calculation or source writes."""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit
from zoneinfo import ZoneInfo

import openpyxl
from jsonschema import Draft7Validator

ROOT = Path(__file__).resolve().parents[1]
DIVISIONS = json.loads((ROOT / "src/divisions.json").read_text())
SCHEMA = json.loads((ROOT / "src/data.schema.json").read_text())
MATCH_ID = re.compile(r"(?:[XYZP]_\d+|[QSF]\d+)$")
SLOT = re.compile(r"(?:[XYZ][1-6]|[1-6][XYZ]|(?:[1-9]|1[0-2])P|[WL][QS]\d+)$")

# Explicit inspected layouts share the same parser; no row/name inference or fuzzy joins.
LAYOUTS = {
    "mixed-a": {
        "team_rows": range(5, 17), "initial": [("X", range(21, 27)), ("Y", range(29, 35))],
        "second": [("high", "POOL ALTO", range(21, 25), range(41, 47)),
                   ("middle", "POOL MEDIO", range(27, 31), range(49, 55)),
                   ("low", "POOL BASSO", range(33, 37), range(57, 63))],
        "standings_header": 19, "game_headers": (4, 21, 40, 48, 56, 66),
        "extra_game_headers": (), "sections": {"J3": "PRIMA FASE A GIRONI", "J39": "SECONDA FASE A GIRONI", "J65": "FASE FINALE"},
        "second_start": 39, "final_start": 65, "final_rows": range(20, 32), "rules_start": 39,
    },
    "mixed-b": {
        "team_rows": range(5, 23), "initial": [("X", range(27, 33)), ("Y", range(35, 41)), ("Z", range(43, 49))],
        "second": [(f"power-{rank}", f"POWER POOL {3 * rank - 2}°-{3 * rank}° POSTO",
                    range(22 + 5 * rank, 25 + 5 * rank), range(53 + 5 * rank, 56 + 5 * rank)) for rank in range(1, 7)],
        "standings_header": 25, "game_headers": (4, 21, 38, 57),
        "extra_game_headers": (62, 67, 72, 77, 82), "sections": {"J3": "PRIMA FASE A GIRONI", "J56": "SECONDA FASE A GIRONI"},
        "second_start": 56, "final_start": None, "final_rows": range(26, 44), "rules_start": 57,
    },
}


class ImportFailure(ValueError):
    """Source no longer satisfies the inspected schema; retain previous JSON."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ImportFailure(message)


class Sheet:
    """Read cached cells and expand only actual merge ranges in memory."""

    def __init__(self, cached: Any, formulas: Any) -> None:
        self.cached, self.formulas = cached, formulas
        self.anchors: dict[tuple[int, int], tuple[int, int]] = {}
        for merged in cached.merged_cells.ranges:
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    self.anchors[row, col] = (merged.min_row, merged.min_col)

    def get(self, row: int, col: int) -> Any:
        row, col = self.anchors.get((row, col), (row, col))
        cell = self.cached.cell(row, col)
        require(cell.data_type != "e", f"{self.cached.title}!{cell.coordinate}: Excel error {cell.value}")
        require(not (self.formulas.cell(row, col).data_type == "f" and cell.value is None),
                f"{self.cached.title}!{cell.coordinate}: formula has no cached result")
        return None if cell.value == "" else cell.value

    def header(self, coordinate: str, expected: str) -> None:
        cell = self.cached[coordinate]
        require(self.get(cell.row, cell.column) == expected,
                f"{self.cached.title}!{coordinate}: expected header {expected!r}")

    def hyperlink(self, row: int, col: int) -> str | None:
        """Preserve the actual hyperlink target, including merged-cell anchors."""
        row, col = self.anchors.get((row, col), (row, col))
        cell = self.cached.cell(row, col)
        if cell.hyperlink is None:
            require(self.get(row, col) is None, f"{self.cached.title}!{cell.coordinate}: link label without hyperlink")
            return None
        return string(cell.hyperlink.target, f"{self.cached.title}!{cell.coordinate}: hyperlink")


def provenance(sheet: str, cell_range: str) -> dict:
    return {"sheet": sheet, "range": cell_range}


def string(value: Any, location: str, *, optional: bool = False) -> str | None:
    if value is None and optional:
        return None
    require(isinstance(value, str) and bool(value.strip()), f"{location}: expected nonempty text")
    return value  # Preserve canonical spelling; no fuzzy joins.


def numeric(value: Any, location: str, *, integer: bool = False) -> int | float | None:
    if value is None:
        return None
    require(type(value) in (int, float) and math.isfinite(value) and value >= 0,
            f"{location}: expected finite nonnegative number, received {value!r}")
    require(not integer or value == int(value), f"{location}: expected integer")
    return int(value) if integer else value


def local_date(value: Any, location: str) -> str | None:
    if value is None:
        return None
    require(isinstance(value, (datetime, date)), f"{location}: expected Excel date")
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()


def local_time(value: Any, location: str) -> str | None:
    if value is None:
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M")
    require(isinstance(value, str) and re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", value) is not None,
            f"{location}: expected HH:mm time")
    return value


def form_team_match(url: str | None, match_id: str, evaluating: str | None,
                    evaluated: str | None) -> bool | None:
    """Validate the organizer's exact Google Form link; flag wrong team prefills."""
    if url is None:
        return None
    try:
        parsed = urlsplit(url)
    except ValueError as exc:
        raise ImportFailure(f"{match_id}: invalid form URL") from exc
    require(parsed.scheme == "https" and parsed.netloc == "docs.google.com"
            and re.fullmatch(r"/forms/d/e/[A-Za-z0-9_-]+/viewform", parsed.path) is not None
            and not parsed.fragment, f"{match_id}: expected an HTTPS Google Form view URL")
    query = parse_qs(parsed.query, keep_blank_values=True)
    require(query.get("entry.1623410474") == [match_id], f"{match_id}: form link references a different match")
    if evaluating is None or evaluated is None:
        return None
    return (query.get("entry.1084977692") == [evaluating]
            and query.get("entry.839559032") == [evaluated])


def parse_workbook(payload: bytes, fetched_at: str, source_mode: str = "live", *, division_id: str = "mixed-a") -> dict:
    """Normalize the inspected four-sheet contract and validate before output."""
    require(division_id in DIVISIONS, f"Unknown division {division_id!r}")
    division, layout = DIVISIONS[division_id], LAYOUTS[division_id]
    try:
        cached = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(payload), data_only=False)
    except Exception as exc:
        raise ImportFailure(f"Not a readable XLSX workbook: {exc}") from exc
    names = ["Classifiche e format", "Partite", "Valutazioni SOTG", "Statistiche"]
    require(all(name in cached.sheetnames for name in names), "Missing required workbook sheets")
    sheets = {name: Sheet(cached[name], formulas[name]) for name in names}
    standings, games, spirit, stats = (sheets[name] for name in names)
    for coordinate, expected in {"B3": f"CAMPIONATI ITALIANI ULTIMATE 2026 - {division['label'].upper()}",
                                 "B4": "Seed", "C4": "Squadra", "D4": "ASD/SSD", "J4": "Provincia",
                                 f"B{layout['standings_header']}": "PRIMA FASE A GIRONI",
                                 f"H{layout['standings_header']}": "SECONDA FASE A GIRONI",
                                 f"N{layout['standings_header']}": "CLASSIFICA FINALE"}.items():
        standings.header(coordinate, expected)
    for pool, rows in layout["initial"]:
        standings.header(f"B{rows.start - 1}", f"GIRONE {pool}")
    for _, label, rows, _ in layout["second"]:
        standings.header(f"H{rows.start - 1}", label)
    for row in layout["game_headers"]:
        for col, expected in {"B": "TAPPA", "C": "LUOGO", "D": "DATA", "E": "ORARIO", "H": "CAMPO", "J": "PARTITE", "R": "SOTG"}.items():
            games.header(f"{col}{row}", expected)
    for row in layout["extra_game_headers"]:
        for col, expected in {"E": "ORARIO", "H": "CAMPO", "J": "PARTITE", "R": "SOTG"}.items():
            games.header(f"{col}{row}", expected)
    for coordinate, expected in layout["sections"].items():
        games.header(coordinate, expected)
    for coordinate, expected in {"A1": "PARTITA", "B1": "DATA", "C1": "SQUADRA VALUTANTE", "D1": "SQUADRA VALUTATA",
                                 "E1": "LINK",
                                 "G1": "Conoscenza e applicazione delle regole", "H1": "Falli e contatto fisico",
                                 "I1": "Comportamento leale e corretto", "J1": "Atteggiamento positivo e autocontrollo",
                                 "K1": "Comunicazione", "L1": "TOTALE", "M1": "COMMENTO"}.items():
        spirit.header(coordinate, expected)
    for coordinate, expected in {"A1": "SQUADRA", "B1": "COGNOME", "C1": "NOME", "D1": "STATISTICHE PER PARTITA   [ assist | mete ]",
                                 "N1": "tot.\nassist", "O1": "tot.\nmete", "P1": "TOTALE"}.items():
        stats.header(coordinate, expected)
    for index in range(10):
        require(stats.get(2, index + 4) == f"{index + 1}a", "Statistiche: numbered match-slot headings changed")

    # Initial slots are explicit identities, distinct from rank qualifiers (1X).
    name_slots = {}
    for pool, rows in layout["initial"]:
        for row in rows:
            slot, name = standings.get(row, 2), standings.get(row, 3)
            require(isinstance(slot, str) and re.fullmatch(f"{pool}[1-6]", slot), f"Initial pool row {row}: invalid slot")
            require(name not in name_slots and slot not in name_slots.values(), "Initial team-name/slot collision")
            name_slots[string(name, f"C{row}")] = slot
    teams = []
    for row in layout["team_rows"]:
        name = string(standings.get(row, 3), f"C{row}")
        require(name in name_slots, f"Team {name!r} absent from initial slots")
        slot = name_slots[name]
        teams.append({"id": slot.lower(), "slot": slot, "name": name,
                      "seed": numeric(standings.get(row, 2), f"B{row}", integer=True),
                      "affiliation": string(standings.get(row, 4), f"D{row}"), "province": string(standings.get(row, 10), f"J{row}"),
                      "pool": slot[0], "provenance": provenance(names[0], f"B{row}:L{row}")})
    by_name = {team["name"]: team["id"] for team in teams}
    by_slot = {team["slot"]: team["id"] for team in teams}

    def participant(slot: Any, name: Any, location: str) -> dict:
        require(isinstance(slot, str) and SLOT.fullmatch(slot) is not None, f"{location}: unknown participant slot {slot!r}")
        if name is not None:
            require(name in by_name, f"{location}: unknown exact team name {name!r}")
        team_id = by_name[name] if name is not None else None
        if slot in by_slot:
            require(team_id == by_slot[slot], f"{location}: initial slot/name mismatch or missing cached team")
        return {"slot": slot, "teamId": team_id}

    def score_at(row: int) -> dict:
        require(games.get(row, 13) in (None, "-", "–", ":"), f"Partite!M{row}: unsupported result/status marker")
        return {"a": numeric(games.get(row, 12), f"L{row}", integer=True),
                "b": numeric(games.get(row, 14), f"N{row}", integer=True)}

    second_pool_by_row = {row: pool for pool, _, _, rows in layout["second"] for row in rows}
    initial_pool_ids = "".join(pool for pool, _ in layout["initial"])
    matches, carryovers = [], []
    for row in range(5, cached[names[1]].max_row + 1):
        match_id, marker = games.get(row, 9), games.get(row, 17)
        require(marker in (None, "report", "ereditata"), f"Partite!Q{row}: unknown report/carry-over marker")
        report_url = None
        if marker == "report":
            require(match_id is not None, f"Partite!Q{row}: report without match ID")
            report_url = games.hyperlink(row, 17)
            require(report_url is not None and re.fullmatch(r"https://docs\.google\.com/spreadsheets/d/[A-Za-z0-9_-]+/edit(?:[?#][^\s]*)?", report_url),
                    f"Partite!Q{row}: expected a public Google Sheets report link")
        if marker == "ereditata":
            require(match_id is None, f"Partite!Q{row}: inherited row has a match ID")
            pool = second_pool_by_row.get(row)
            require(pool is not None, "Carry-over outside second phase")
            carryovers.append({"pool": pool,
                               "a": participant(games.get(row, 10), games.get(row, 11), f"J{row}"),
                               "b": participant(games.get(row, 16), games.get(row, 15), f"P{row}"),
                               "score": score_at(row), "provenance": provenance(names[1], f"J{row}:Q{row}")})
        if match_id is None:
            # Never silently drop a fixture with a removed ID.
            require(marker is not None or not (isinstance(games.get(row, 10), str) and SLOT.fullmatch(games.get(row, 10))),
                    f"Partite row {row}: participants present without match ID")
            continue
        require(isinstance(match_id, str) and MATCH_ID.fullmatch(match_id) is not None, f"Partite!I{row}: invalid match ID")
        if row < layout["second_start"]:
            phase, pool = "first", match_id[0]
        elif layout["final_start"] is not None and row >= layout["final_start"]:
            phase, pool = "finals", None
        else:
            phase, pool = "second", second_pool_by_row.get(row)
        require(match_id[0] in {"first": initial_pool_ids, "second": "P", "finals": "QSF"}[phase], f"{match_id}: phase/ID mismatch")
        require(phase != "second" or pool is not None, f"{match_id}: fixture outside a verified second-phase block")
        score = score_at(row)
        state = "none" if all(value is None for value in score.values()) else "complete" if all(value is not None for value in score.values()) else "partial"
        matches.append({"id": match_id, "phase": phase, "pool": pool,
                        "stage": numeric(games.get(row, 2), f"B{row}", integer=True),
                        "city": string(games.get(row, 3), f"C{row}", optional=True),
                        "date": local_date(games.get(row, 4), f"D{row}"),
                        "start": local_time(games.get(row, 5), f"E{row}"), "end": local_time(games.get(row, 7), f"G{row}"),
                        "field": string(games.get(row, 8), f"H{row}", optional=True),
                        "a": participant(games.get(row, 10), games.get(row, 11), f"J{row}"),
                        "b": participant(games.get(row, 16), games.get(row, 15), f"P{row}"),
                        "score": score, "resultState": state, "reportUrl": report_url,
                        "provenance": provenance(names[1], f"B{row}:U{row}")})

    pools = []
    pool_specs = [(pool, f"Girone {pool}", rows, 2, "first") for pool, rows in layout["initial"]]
    pool_specs += [(pool, label.capitalize(), rows, 8, "second") for pool, label, rows, _ in layout["second"]]
    for pool_id, label, rows, offset, phase in pool_specs:
        entries = []
        for row in rows:
            entries.append({**participant(standings.get(row, offset), standings.get(row, offset + 1), f"pool {pool_id} row {row}"),
                            "points": numeric(standings.get(row, offset + 2), f"pool {pool_id} points"),
                            "goalsFor": numeric(standings.get(row, offset + 3), "goals for", integer=True),
                            "goalsAgainst": numeric(standings.get(row, offset + 4), "goals against", integer=True),
                            "provenance": provenance(names[0], f"{openpyxl.utils.get_column_letter(offset)}{row}:{openpyxl.utils.get_column_letter(offset + 4)}{row}")})
        pools.append({"id": pool_id, "name": label, "phase": phase, "entries": entries,
                      "provenance": provenance(names[0], f"{openpyxl.utils.get_column_letter(offset)}{rows.start}:{openpyxl.utils.get_column_letter(offset + 4)}{rows.stop - 1}")})
    final_classification = []
    for row in layout["final_rows"]:
        position = row - layout["final_rows"].start + 1
        require(standings.get(row, 14) == position, f"Final position N{row} changed")
        name = standings.get(row, 15)
        if name is not None:
            require(name in by_name, f"O{row}: unknown final team")
            final_classification.append({"position": position, "teamId": by_name[name], "provenance": provenance(names[0], f"N{row}:O{row}")})

    sotg = []
    for row in range(2, cached[names[2]].max_row + 1):
        match_id = spirit.get(row, 1)
        if match_id is None:
            require(all(spirit.get(row, col) is None for col in (3, 4, 5, 7, 8, 9, 10, 11, 12, 13)), f"SOTG row {row}: evaluation without match ID")
            continue
        evaluating, evaluated = spirit.get(row, 3), spirit.get(row, 4)
        require(evaluating is None or evaluating in by_name, f"SOTG row {row}: unknown evaluating team")
        require(evaluated is None or evaluated in by_name, f"SOTG row {row}: unknown evaluated team")
        categories = [numeric(spirit.get(row, col), f"SOTG row {row} category", integer=True) for col in range(7, 12)]
        total = numeric(spirit.get(row, 12), f"SOTG row {row} total")
        comment = string(spirit.get(row, 13), f"SOTG row {row} comment", optional=True)
        if any(value is not None for value in [*categories, total, comment]):
            require(evaluating is not None and evaluated is not None, f"SOTG row {row}: published evaluation needs both teams")
        if total is not None and all(value is not None for value in categories):
            require(sum(categories) == total, f"SOTG row {row}: total disagrees with category sum")
        form_url = spirit.hyperlink(row, 5)
        sotg.append({"matchId": match_id, "date": local_date(spirit.get(row, 2), f"SOTG B{row}"),
                     "evaluatingTeamId": by_name.get(evaluating), "evaluatedTeamId": by_name.get(evaluated),
                     "categories": categories, "total": total, "comment": comment,
                     "formUrl": form_url,
                     "formTeamsMatch": form_team_match(form_url, match_id, evaluating, evaluated),
                     "provenance": provenance(names[2], f"A{row}:M{row}")})
    players = []
    for row in range(3, cached[names[3]].max_row + 1):
        values = [stats.get(row, col) for col in range(1, 17)]
        if all(value is None for value in values):
            continue
        require(values[0] in by_name, f"Statistiche row {row}: unknown or missing team")
        # The organizer can publish totals before filling a player's name.
        # Preserve the row and missing identity instead of dropping its statistics.
        players.append({"teamId": by_name[values[0]], "surname": string(values[1], f"B{row}", optional=True),
                        "givenName": string(values[2], f"C{row}", optional=True),
                        "assists": numeric(values[13], f"N{row}", integer=True), "goals": numeric(values[14], f"O{row}", integer=True),
                        "total": numeric(values[15], f"P{row}", integer=True),
                        "matchSlots": [None if value is None else str(value) for value in values[3:13]],
                        "provenance": provenance(names[3], f"A{row}:P{row}")})
    slots = {}
    for side in [side for match in matches for side in (match["a"], match["b"])] + [entry for pool in pools for entry in pool["entries"]]:
        if side["slot"] not in by_slot:
            require(side["slot"] not in slots or slots[side["slot"]] == side["teamId"], f"Conflicting resolution of slot {side['slot']}")
            slots[side["slot"]] = side["teamId"]
    reported = games.get(1, 4)
    if reported is not None:
        try:
            reported = datetime.strptime(reported, "%d/%m/%Y %H:%M:%S").replace(tzinfo=ZoneInfo("Europe/Rome")).isoformat()
        except (ValueError, TypeError) as exc:
            raise ImportFailure("Partite!D1: source timestamp format changed") from exc
    rules = [{"text": str(standings.get(row, col)).strip(), "provenance": provenance(names[0], f"{openpyxl.utils.get_column_letter(col)}{row}")}
             for col in (2, 14) for row in range(layout["rules_start"], cached[names[0]].max_row + 1) if standings.get(row, col) is not None]
    data = {"schemaVersion": 1, "division": division_id, "sourceUrl": division["sourceUrl"], "sourceMode": source_mode, "sourceReportedAt": reported,
            "fetchedAt": fetched_at, "workbookSha256": hashlib.sha256(payload).hexdigest(), "teams": teams, "matches": matches,
            "pools": pools, "phaseSlots": [{"slot": slot, "teamId": team_id} for slot, team_id in slots.items()],
            "carryovers": carryovers, "finalClassification": final_classification,
            "sotgCategories": [spirit.get(1, col) for col in range(7, 12)], "sotg": sotg, "players": players, "rules": rules,
            "availability": {"sotg": any(any(value is not None for value in [*evaluation["categories"], evaluation["total"], evaluation["comment"]]) for evaluation in sotg),
                             "teamPlayers": bool(players), "matchPlayers": False, "roster": False, "sotgSummary": False}}
    validate_data(data)
    cached.close()
    formulas.close()
    return data


def validate_data(data: dict) -> None:
    """Check JSON shape plus referential and semantic invariants once."""
    errors = sorted(Draft7Validator(SCHEMA).iter_errors(data), key=lambda error: str(error.path))
    require(not errors, f"JSON schema: {errors[0].message if errors else ''}")
    division = DIVISIONS.get(data.get("division", "mixed-a"))
    require(division is not None and data["sourceUrl"] == division["sourceUrl"], "Division/source mismatch")
    for field in ("fetchedAt", "sourceReportedAt"):
        if data[field] is not None:
            try:
                stamp = datetime.fromisoformat(data[field].replace("Z", "+00:00"))
                require(stamp.utcoffset() is not None, f"{field}: timezone required")
            except ValueError as exc:
                raise ImportFailure(f"{field}: invalid timestamp") from exc
    teams = {team["id"] for team in data["teams"]}
    require(len(teams) == len(data["teams"]) and bool(teams), "Duplicate or empty teams")
    require(len({team['name'] for team in data['teams']}) == len(teams), "Duplicate team names")
    require(len({team['seed'] for team in data['teams']}) == len(teams), "Duplicate seeds")
    matches = {match["id"]: match for match in data["matches"]}
    require(len(matches) == len(data["matches"]) and bool(matches), "Duplicate or empty match IDs")
    for match in data["matches"]:
        if match["date"]:
            try:
                date.fromisoformat(match["date"])
            except ValueError as exc:
                raise ImportFailure(f"{match['id']}: invalid calendar date") from exc
        score = match["score"]
        state = "none" if score["a"] is None and score["b"] is None else "complete" if score["a"] is not None and score["b"] is not None else "partial"
        require(match["resultState"] == state, f"{match['id']}: inconsistent result state")
        require(state != "complete" or (match["a"]["teamId"] is not None and match["b"]["teamId"] is not None), f"{match['id']}: completed result without resolved teams")
        require(match["a"]["teamId"] is None or match["a"]["teamId"] != match["b"]["teamId"], f"{match['id']}: team plays itself")
    sides = [side for match in data["matches"] for side in (match["a"], match["b"])]
    sides += [entry for pool in data["pools"] for entry in pool["entries"]] + data["phaseSlots"]
    sides += [side for carry in data["carryovers"] for side in (carry["a"], carry["b"])]
    require(all(side["teamId"] is None or side["teamId"] in teams for side in sides), "Unknown team reference")
    directions = set()
    names = {team["id"]: team["name"] for team in data["teams"]}
    for evaluation in data["sotg"]:
        require(evaluation["matchId"] in matches, "SOTG references unknown match ID")
        match = matches[evaluation["matchId"]]
        require(evaluation["date"] == match["date"], "SOTG date disagrees with match")
        evaluating, evaluated = evaluation["evaluatingTeamId"], evaluation["evaluatedTeamId"]
        expected_form_match = form_team_match(evaluation.get("formUrl"), evaluation["matchId"],
                                             names.get(evaluating), names.get(evaluated))
        require(evaluation.get("formTeamsMatch") == expected_form_match, "Inconsistent form team-prefill flag")
        if evaluating is not None or evaluated is not None:
            require(evaluating is not None and evaluated is not None and evaluating != evaluated,
                    "SOTG evaluation requires two opposite teams")
            require({evaluating, evaluated} == {match["a"]["teamId"], match["b"]["teamId"]}, "SOTG direction does not match participants")
            key = (evaluation["matchId"], evaluating, evaluated)
            require(key not in directions, "Duplicate SOTG evaluation direction")
            directions.add(key)
    require(all(player["teamId"] in teams for player in data["players"]), "Unknown player team")
    require(all(entry["teamId"] in teams for entry in data["finalClassification"]), "Unknown final team")
    require(len({entry["teamId"] for entry in data["finalClassification"]}) == len(data["finalClassification"]), "Duplicate final team")
